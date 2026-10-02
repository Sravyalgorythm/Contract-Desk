from io import BytesIO
from unittest.mock import patch

from docx import Document

from legal_agent.documents import DocumentBlock, SourceRef, extract_blocks
from legal_agent.retrieval import chunk_blocks


def test_pdf_extraction_keeps_page_numbers():
    pages = [type("Page", (), {"extract_text": lambda self: "Payment terms"})() for _ in range(2)]
    reader = type("Reader", (), {"is_encrypted": False, "pages": pages})
    with patch("legal_agent.documents.PdfReader", return_value=reader):
        blocks = extract_blocks(b"pdf bytes", "agreement.pdf")

    assert [block.source.locator for block in blocks] == ["PDF page 1", "PDF page 2"]


def test_docx_extraction_tracks_headings_paragraphs_and_tables():
    document = Document()
    document.add_heading("Payment", level=1)
    document.add_paragraph("Invoices are due in thirty days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Late fee"
    table.cell(0, 1).text = "1 percent"
    output = BytesIO()
    document.save(output)

    blocks = extract_blocks(output.getvalue(), "agreement.docx")

    assert "Payment / paragraph 2" == blocks[1].source.locator
    assert blocks[2].source.locator == "Payment / table 1, row 1"
    assert "Late fee | 1 percent" == blocks[2].text


def test_chunking_keeps_source_blocks_separate():
    blocks = [
        DocumentBlock("Payment is due monthly.", SourceRef("contract.pdf", "PDF page 1")),
        DocumentBlock("Renewal is automatic.", SourceRef("contract.pdf", "PDF page 2")),
    ]

    chunks = chunk_blocks(blocks, max_words=10, overlap_words=2)

    assert len(chunks) == 2
    assert chunks[0].sources == (blocks[0].source,)
    assert chunks[1].sources == (blocks[1].source,)


def test_oversized_block_splits_without_losing_source():
    block = DocumentBlock(
        "one two three four five six seven eight nine ten",
        SourceRef("contract.pdf", "PDF page 3"),
    )

    chunks = chunk_blocks([block], max_words=4, overlap_words=1)

    assert [chunk.text for chunk in chunks] == [
        "one two three four",
        "four five six seven",
        "seven eight nine ten",
    ]
    assert all(chunk.sources == (block.source,) for chunk in chunks)
